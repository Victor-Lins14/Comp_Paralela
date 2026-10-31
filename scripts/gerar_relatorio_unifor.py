# -*- coding: utf-8 -*-
"""
Gera o relatório acadêmico a partir dos CSVs do benchmark, em dois formatos:

1. Relatorio_Analise_Desempenho_UNIFOR.docx — montado SOBRE o arquivo
   templates/Modelo-Artigo2.docx (modelo oficial fornecido pelo professor), preservando
   cabeçalho/logo, margens e as caixas de título de seção do modelo.
2. RELATORIO_UNIFOR.md — mesma redação em Markdown (para leitura no GitHub).

Todos os números citados no texto são lidos dos CSVs em resultados/, então basta rodar
novamente o benchmark (java ... Main --cli), depois scripts/gerar_graficos.py e por fim
este script para o relatório refletir os novos dados.

Dependências: pandas, python-docx  (pip install pandas python-docx)
"""

import copy
import datetime
import os
import sys

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor

# ---------------------------------------------------------------------------
# Configuração (edite aqui)
# ---------------------------------------------------------------------------
AUTOR_1 = "Victor Lins"
AUTOR_2 = "Diogo Guimarães Gifoni"
# Máquina onde o benchmark foi executado. Se PROCESSADOR ficar vazio, usa o que o Java
# registrou em resultados/ambiente.txt. NUCLEOS_FISICOS = 0 omite a discussão sobre SMT.
PROCESSADOR = "AMD Ryzen 7 7735HS, com 8 núcleos físicos e 16 threads via SMT"
MAQUINA = "notebook Acer Nitro V"
NUCLEOS_FISICOS = 8
URL_GITHUB = "https://github.com/Victor-Lins14/Comp_Paralela"
TITULO = ("Análise de Desempenho de Algoritmos de Busca em Ambientes Concorrentes e Paralelos: "
          "Um Estudo Comparativo em Java")
PALAVRAS_CHAVE = ("Computação Paralela. Algoritmos de Ordenação. Java Fork/Join. "
                  "Speedup. Análise Estatística.")

TEMPLATE = "templates/Modelo-Artigo2.docx"
SAIDA_DOCX = "Relatorio_Analise_Desempenho_UNIFOR.docx"
SAIDA_MD = "RELATORIO_UNIFOR.md"
DIR_FIG = "resultados/graficos_relatorio"

MESES = ["jan.", "fev.", "mar.", "abr.", "maio", "jun.", "jul.", "ago.", "set.", "out.", "nov.", "dez."]
HOJE = datetime.date.today()
ANO = HOJE.year
FONTE = f"Fonte: elaborado pelos autores ({ANO})."


# ---------------------------------------------------------------------------
# Formatação numérica pt-BR
# ---------------------------------------------------------------------------
def num(x, casas=2):
    s = f"{x:,.{casas}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def ms(x):
    return f"{num(x, 2)} ms"


def sp(x):
    return f"{num(x, 2)}x"


def pct(x):
    return f"{num(x * 100.0, 1)}%"


def inteiro(x):
    return num(x, 0)


# ---------------------------------------------------------------------------
# Leitura dos dados
# ---------------------------------------------------------------------------
def carregar_ambiente():
    amb = {}
    caminho = "resultados/ambiente.txt"
    if os.path.exists(caminho):
        with open(caminho, encoding="utf-8") as f:
            for linha in f:
                if "=" in linha:
                    k, v = linha.strip().split("=", 1)
                    amb[k] = v
    return amb


class Dados:
    def __init__(self):
        self.resumo = pd.read_csv("resultados/benchmark_resumo.csv")
        self.intervalos = pd.read_csv("resultados/benchmark_intervalos.csv")
        self.amb = carregar_ambiente()

    def m(self, alg, tam, thr, col="MediaMs", nat="ALEATORIO"):
        r = self.resumo
        f = r[(r.Algoritmo == alg) & (r.Tamanho == tam) & (r.Threads == thr) & (r.NaturezaDados == nat)]
        return float(f[col].values[0]) if not f.empty else float("nan")

    def i(self, alg, k, thr, col="MediaMs"):
        r = self.intervalos
        f = r[(r.Algoritmo == alg) & (r.IntervaloMax == k) & (r.Threads == thr)]
        return float(f[col].values[0]) if not f.empty else float("nan")

    def melhor(self, alg, tam, nat="ALEATORIO"):
        """(threads, speedup) da configuração paralela com maior Speedup."""
        r = self.resumo
        f = r[(r.Algoritmo == alg) & (r.Tamanho == tam) & (r.NaturezaDados == nat) & (r.Threads > 1)]
        if f.empty:
            return 1, float("nan")
        linha = f.loc[f.Speedup.idxmax()]
        return int(linha.Threads), float(linha.Speedup)

    def threads(self):
        return sorted(int(t) for t in self.resumo.Threads.unique())

    def linhas_tabela(self, tam, algs):
        linhas = []
        for alg in algs:
            f = self.resumo[(self.resumo.Algoritmo == alg) & (self.resumo.Tamanho == tam) &
                            (self.resumo.NaturezaDados == "ALEATORIO")].sort_values("Threads")
            for _, r in f.iterrows():
                cv = r.DesvioPadraoMs / r.MediaMs if r.MediaMs > 0 else 0.0
                linhas.append([alg, r.Modo, str(int(r.Threads)), num(r.MediaMs, 3), num(r.DesvioPadraoMs, 3),
                               num(r.MedianaMs, 3), pct(cv), sp(r.Speedup), pct(r.Eficiencia)])
        return linhas

    def linhas_intervalo(self):
        linhas = []
        for _, r in self.intervalos.sort_values(["IntervaloMax", "Algoritmo", "Threads"]).iterrows():
            linhas.append([r.Algoritmo, r.Modo, str(int(r.Threads)), inteiro(r.IntervaloMax),
                           num(r.MediaMs, 3), num(r.DesvioPadraoMs, 3), sp(r.Speedup)])
        return linhas


# ---------------------------------------------------------------------------
# Conteúdo (lista de blocos renderizada em DOCX e em Markdown)
# ---------------------------------------------------------------------------
def montar_conteudo(d: Dados):
    B = []
    thr = d.threads()
    p_max = max(thr)
    lista_thr = ", ".join(str(t) for t in thr if t > 1)
    lista_thr = lista_thr[::-1].replace(" ,", " e ", 1)[::-1]

    # Ambiente
    n_logicos = d.amb.get("processadores_logicos", "?")
    cpu = PROCESSADOR or d.amb.get("processador", "processador não identificado")
    so = d.amb.get("sistema_operacional", "sistema operacional não identificado")
    jv = d.amb.get("java_version", "?")
    jvm = d.amb.get("java_vm", "Java HotSpot VM")

    # Métricas-chave
    t_ms_1m, s_ms_1m = d.melhor("Merge Sort", 1_000_000)
    t_qs_1m, s_qs_1m = d.melhor("Quick Sort", 1_000_000)
    t_cs_1m, s_cs_1m = d.melhor("Counting Sort", 1_000_000)
    t_bs, s_bs = d.melhor("Bubble Sort", 50_000)
    ef_bs_max = d.m("Bubble Sort", 50_000, p_max, "Eficiencia")
    ef_ms_max = d.m("Merge Sort", 1_000_000, p_max, "Eficiencia")
    ef_qs_max = d.m("Quick Sort", 1_000_000, p_max, "Eficiencia")
    s_qs_10k = d.m("Quick Sort", 10_000, p_max, "Speedup")
    s_ms_10k = d.m("Merge Sort", 10_000, p_max, "Speedup")

    cs_k_pequeno_s = d.i("Counting Sort", 100_000, 1)
    cs_k_grande_s = d.i("Counting Sort", 10_000_000, 1)
    cs_k_grande_p = d.i("Counting Sort", 10_000_000, 8)
    qs_k_grande_s = d.i("Quick Sort", 10_000_000, 1)

    # SMT: threads acima do número de núcleos físicos compartilham as unidades de execução
    texto_smt = ""
    if NUCLEOS_FISICOS and p_max > NUCLEOS_FISICOS:
        s8_ms = d.m("Merge Sort", 1_000_000, NUCLEOS_FISICOS, "Speedup")
        s8_qs = d.m("Quick Sort", 1_000_000, NUCLEOS_FISICOS, "Speedup")
        smax_ms = d.m("Merge Sort", 1_000_000, p_max, "Speedup")
        smax_qs = d.m("Quick Sort", 1_000_000, p_max, "Speedup")
        texto_smt = (
            f" Há ainda um limite de hardware: o processador possui {NUCLEOS_FISICOS} núcleos físicos, de modo que "
            f"acima de {NUCLEOS_FISICOS} threads os núcleos passam a ser compartilhados por duas threads via SMT "
            "(Simultaneous Multithreading), que dividem as mesmas unidades de execução e caches. Por isso, de "
            f"{NUCLEOS_FISICOS} para {p_max} threads o Speedup passou de {sp(s8_ms)} para {sp(smax_ms)} no Merge "
            f"Sort e de {sp(s8_qs)} para {sp(smax_qs)} no Quick Sort")
        if smax_ms < s8_ms or smax_qs < s8_qs:
            texto_smt += ", chegando a cair: o custo de coordenar mais threads superou o ganho do SMT."
        elif smax_ms < 1.5 * s8_ms and smax_qs < 1.5 * s8_qs:
            texto_smt += f", ganho bem menor que o obtido até {NUCLEOS_FISICOS} threads, quando cada thread tinha um núcleo próprio."
        else:
            texto_smt += f", com eficiência por thread menor que a obtida até {NUCLEOS_FISICOS} threads."

    # Speedup superlinear (Sp > p) no Bubble Sort: efeito de cache, citado no texto se ocorrer
    sup = d.resumo[(d.resumo.Algoritmo == "Bubble Sort") & (d.resumo.Tamanho == 50_000) &
                   (d.resumo.NaturezaDados == "ALEATORIO") & (d.resumo.Threads > 1) &
                   (d.resumo.Speedup > d.resumo.Threads)]
    superlinear_bs = ""
    if not sup.empty:
        p_sup = ", ".join(str(int(t)) for t in sorted(sup.Threads))
        superlinear_bs = (
            f" Com {p_sup} thread(s) observou-se Speedup superlinear (Sp > p). Como o trabalho total é idêntico ao "
            "da versão serial, o efeito pode ser explicado pela hierarquia de memória: cada fase percorre o vetor "
            "inteiro (200 KB para N = 50.000), mas, dividido entre p threads, o trecho de cada uma fica pequeno o "
            "suficiente para permanecer na cache privada do seu núcleo, reduzindo o custo de cada acesso.")

    # --- Cabeçalho -------------------------------------------------------
    B.append(("titulo", TITULO))
    B.append(("autor", f"Autor 1: {AUTOR_1}."))
    B.append(("autor", f"Autor 2: {AUTOR_2}."))
    B.append(("palavras", f"Palavras-chave: {PALAVRAS_CHAVE}"))

    # --- Resumo ----------------------------------------------------------
    B.append(("secao", "Resumo"))
    B.append(("par",
        "Este trabalho analisa experimentalmente o desempenho de quatro algoritmos de ordenação, etapa que "
        "antecede e viabiliza a busca eficiente de dados, em versões sequenciais e paralelas implementadas em "
        "Java: Quick Sort, Merge Sort, Counting Sort e Bubble Sort (na formulação paralela Odd-Even "
        "Transposition). Foi desenvolvido um framework de testes que varia o tamanho dos conjuntos de dados "
        "(10.000 a 1.000.000 de elementos), a natureza da entrada (aleatória, inversamente ordenada, quase "
        f"ordenada e com poucas chaves distintas), o intervalo de valores e o número de threads (1, {lista_thr}), "
        "com cinco amostras por configuração e registro dos resultados em arquivos CSV. Para N = 1.000.000, "
        f"o Merge Sort e o Quick Sort paralelos atingiram Speedups de {sp(s_ms_1m)} e {sp(s_qs_1m)}, "
        f"respectivamente. O Bubble Sort Odd-Even alcançou {sp(s_bs)} com {t_bs} threads para N = 50.000, "
        "limitado pelo custo das barreiras de sincronização entre fases. O Counting Sort obteve os menores tempos "
        "absolutos, mas, por ser limitado pela memória, praticamente não se beneficiou do paralelismo e mostrou "
        "forte sensibilidade à amplitude k do intervalo de valores. Os resultados foram visualizados em uma "
        "aplicação de gráficos dinâmicos em Java Swing e em Jupyter Notebook."))

    # --- Introdução ------------------------------------------------------
    B.append(("secao", "Introdução"))
    B.append(("par",
        "A ordenação é uma das operações mais estudadas da Ciência da Computação, pois algoritmos de busca "
        "eficientes, como a busca binária em O(log n), as consultas por intervalo e a indexação em bancos de "
        "dados, pressupõem dados previamente ordenados. Com a consolidação dos processadores multicore, "
        "reduzir o tempo dessas operações passou a depender não só da complexidade assintótica, mas também da "
        "capacidade de dividir o trabalho entre núcleos e dos custos de sincronização, comunicação e acesso à "
        "memória que essa divisão introduz."))
    B.append(("par",
        "Dentre os seis algoritmos propostos no enunciado, foram escolhidos quatro que representam paradigmas e "
        "classes de complexidade distintos: (i) Quick Sort, divisão e conquista in-place com complexidade média "
        "O(n log n), aqui com pivô por mediana de três e particionamento de Hoare; (ii) Merge Sort, divisão e "
        "conquista estável com O(n log n) garantido no pior caso; (iii) Counting Sort, método não comparativo "
        "de complexidade O(n + k), em que k é a amplitude do intervalo de valores; e (iv) Bubble Sort, método "
        "elementar O(n²) baseado em trocas de elementos adjacentes."))
    B.append(("par",
        "A abordagem de paralelização utilizou as ferramentas do pacote java.util.concurrent. Quick Sort e "
        "Merge Sort foram paralelizados com o framework Fork/Join (ForkJoinPool e RecursiveAction), cujo "
        "escalonador por roubo de trabalho (work-stealing) distribui as subtarefas recursivas entre as threads, "
        "com um limiar de granularidade abaixo do qual a recursão prossegue sequencialmente. O Counting Sort "
        "paralelo usa histogramas locais por thread (sem locks), redução paralela por faixas de valores e escrita "
        "paralela em regiões disjuntas, limitando o número de histogramas quando k é grande em relação a n. O "
        "Bubble Sort foi paralelizado pelo Odd-Even Transposition Sort, sua formulação paralela clássica: as "
        "comparações de cada fase envolvem pares disjuntos e podem ser feitas simultaneamente, e as fases são "
        "separadas por uma barreira (Phaser). Para que o Speedup meça apenas o ganho do paralelismo, a versão "
        "serial de cada algoritmo executa exatamente o mesmo procedimento da versão paralela em uma única "
        "thread."))

    # --- Metodologia -----------------------------------------------------
    B.append(("secao", "Metodologia"))
    B.append(("par",
        "Foi desenvolvido em Java um framework de testes composto pelas classes ExecutorBenchmark (orquestração "
        "dos experimentos), GeradorDados (geração determinística das entradas) e ExportadorCSV (gravação dos "
        "resultados), além da interface comum AlgoritmoOrdenacao, que obriga cada algoritmo a oferecer um método "
        "serial e um paralelo parametrizado pelo número de threads. Os experimentos foram executados em um "
        f"{MAQUINA or 'computador'} com processador {cpu} ({n_logicos} processadores lógicos visíveis à JVM), "
        f"sistema operacional {so} e "
        f"{jvm} (Java {jv}). Antes das medições, o framework executa ciclos de aquecimento da JVM para que o "
        "compilador Just-In-Time otimize os métodos medidos."))
    B.append(("par",
        "O protocolo seguiu um projeto fatorial com quatro fatores: (1) número de threads: execução serial "
        f"(1 thread) e paralela com {lista_thr} threads, sendo que o paralelismo do pool limita a quantidade de "
        "núcleos usados simultaneamente, o que permite investigar o efeito de disponibilizar mais ou menos "
        "núcleos ao algoritmo; (2) tamanho N: 10.000, 25.000, 50.000, 100.000, 500.000 e 1.000.000 de inteiros, "
        "com o Bubble Sort limitado a N = 50.000 por ser quadrático; (3) natureza dos dados: aleatório uniforme, "
        "inversamente ordenado, quase ordenado (10% dos elementos fora de posição) e poucos valores únicos (até "
        "32 chaves); e (4) intervalo de valores [0..k], com k em {1.000; 100.000; 1.000.000; 10.000.000} para "
        "N = 500.000, comparando-se a execução serial com a paralela de 8 threads."))
    B.append(("par",
        "Para cada configuração foram coletadas 5 amostras, com sementes pseudoaleatórias fixas, de modo que "
        "todos os algoritmos e quantidades de threads processem exatamente os mesmos vetores. O tempo foi medido "
        "com System.nanoTime() e cada resultado foi verificado quanto à ordenação correta. A análise estatística "
        "dos resultados obtidos calcula, por configuração, média, desvio padrão amostral (s), mediana, mínimo, "
        "máximo e coeficiente de variação (CV = s / média), além das métricas de desempenho paralelo Speedup "
        "(Sp = T1 / Tp) e Eficiência (Ep = Sp / p), interpretadas à luz da Lei de Amdahl, para identificar "
        "padrões de desempenho e comparar os algoritmos sob diferentes condições. Os dados foram gravados em três "
        "arquivos CSV (amostras individuais, resumo estatístico e experimento de intervalo de valores) e "
        "visualizados em gráficos gerados em Python, em Jupyter Notebook e em uma aplicação de gráficos dinâmicos "
        "em Java Swing."))

    # --- Resultados ------------------------------------------------------
    B.append(("secao", "Resultados e Discussão"))

    B.append(("sub", "Escalabilidade em função do tamanho dos dados"))
    B.append(("par",
        "A Figura 1 mostra o crescimento do tempo de execução com N. As três classes de complexidade aparecem "
        "separadas por ordens de grandeza: para N = 50.000 elementos aleatórios, a execução serial levou em média "
        f"{ms(d.m('Bubble Sort', 50000, 1))} no Bubble Sort, {ms(d.m('Merge Sort', 50000, 1))} no Merge Sort, "
        f"{ms(d.m('Quick Sort', 50000, 1))} no Quick Sort e {ms(d.m('Counting Sort', 50000, 1))} no Counting Sort. "
        "A Tabela 1 detalha a estatística das cinco amostras para esse tamanho."))
    B.append(("fig", os.path.join(DIR_FIG, "fig1_escalabilidade_tamanho.png"),
              "Tempo médio de execução em função do tamanho N (serial e paralelo)"))
    cab = ["Algoritmo", "Modo", "Threads", "Média (ms)", "Desvio (ms)", "Mediana (ms)", "CV", "Speedup", "Eficiência"]
    B.append(("tab", "Estatísticas para N = 50.000 (dados aleatórios, 5 amostras)", cab,
              d.linhas_tabela(50_000, ["Quick Sort", "Merge Sort", "Counting Sort", "Bubble Sort"])))

    B.append(("sub", "Speedup e eficiência por número de threads"))
    B.append(("par",
        "Para N = 1.000.000 (Tabela 2 e Figuras 2 e 3), os algoritmos de divisão e conquista foram os que mais se "
        f"beneficiaram do paralelismo. O Merge Sort caiu de {ms(d.m('Merge Sort', 1_000_000, 1))} (serial) para "
        f"{ms(d.m('Merge Sort', 1_000_000, t_ms_1m))} com {t_ms_1m} threads (Speedup de {sp(s_ms_1m)}), e o Quick "
        f"Sort caiu de {ms(d.m('Quick Sort', 1_000_000, 1))} para {ms(d.m('Quick Sort', 1_000_000, t_qs_1m))} com "
        f"{t_qs_1m} threads (Speedup de {sp(s_qs_1m)}). Ainda assim, a eficiência com {p_max} threads ficou em "
        f"{pct(ef_ms_max)} e {pct(ef_qs_max)}, respectivamente: como prevê a Lei de Amdahl, as partes "
        "sequenciais (o primeiro particionamento do Quick Sort e a última intercalação do Merge Sort percorrem o "
        "vetor inteiro em uma única thread) e a disputa pela largura de banda da memória limitam o ganho."
        + texto_smt))
    B.append(("par",
        f"Para cargas pequenas o comportamento se inverte: com N = 10.000 e {p_max} threads o Speedup foi de "
        f"{sp(s_qs_10k)} no Quick Sort e {sp(s_ms_10k)} no Merge Sort, pois o custo de criar o pool e distribuir "
        "as tarefas é comparável ao próprio tempo de ordenação. O Counting Sort, por sua vez, já é muito rápido "
        f"na versão serial ({ms(d.m('Counting Sort', 1_000_000, 1))} para N = 1.000.000) e seu melhor Speedup foi "
        f"de {sp(s_cs_1m)} ({t_cs_1m} threads): por ser limitado pela memória, e não pelo processamento, dividir "
        "os acessos entre núcleos não reduz o gargalo, e as fases extras de alocação e redução dos histogramas "
        "locais anulam boa parte do ganho."))
    B.append(("par",
        f"O Bubble Sort Odd-Even (N = 50.000) reduziu o tempo de {ms(d.m('Bubble Sort', 50000, 1))} para "
        f"{ms(d.m('Bubble Sort', 50000, t_bs))} com {t_bs} threads (Speedup de {sp(s_bs)}), com eficiência de "
        f"{pct(ef_bs_max)} em {p_max} threads. Como as versões serial e paralela executam as mesmas fases e "
        "comparações, esse ganho decorre exclusivamente da divisão de cada fase entre as threads. O limite vem "
        "da sincronização: o algoritmo executa até N fases, cada uma terminada por uma barreira; com mais threads, "
        "o trabalho por fase de cada thread diminui (N / 2p comparações), enquanto o custo de cada barreira "
        "aumenta, de modo que a espera passa a dominar o tempo total." + superlinear_bs))
    B.append(("fig", os.path.join(DIR_FIG, "fig2_tempo_por_threads.png"),
              "Tempo médio e desvio padrão (5 amostras) por número de threads"))
    B.append(("fig", os.path.join(DIR_FIG, "fig3_speedup_eficiencia.png"),
              "Speedup (Sp) e Eficiência (Ep) em função do número de threads"))
    B.append(("tab", "Estatísticas para N = 1.000.000 (dados aleatórios, 5 amostras)", cab,
              d.linhas_tabela(1_000_000, ["Quick Sort", "Merge Sort", "Counting Sort"])))

    B.append(("sub", "Impacto da natureza dos dados e do intervalo de valores"))
    B.append(("par",
        "A Figura 4 compara as quatro naturezas de entrada para N = 100.000. Com pivô por mediana de três e "
        "particionamento de Hoare, o Quick Sort manteve o comportamento O(n log n) mesmo nas entradas "
        "inversamente ordenadas e com muitas chaves repetidas, que levariam uma implementação com pivô fixo ao "
        "pior caso O(n²). O Merge Sort foi mais rápido nas entradas quase ordenadas e com poucas chaves "
        "distintas, pois verifica se dois subvetores adjacentes já estão em ordem antes de intercalá-los."))
    B.append(("par",
        "A Figura 5 e a Tabela 3 mostram o tempo em função do intervalo de valores [0..k] para N = 500.000. Quick "
        "Sort e Merge Sort, baseados em comparações, praticamente não variam com k. Já o Counting Sort depende "
        f"diretamente de k: com k = 100.000 levou {ms(cs_k_pequeno_s)} (serial), mas com k = 10.000.000 subiu para "
        f"{ms(cs_k_grande_s)}, pois o vetor de contagem passa a ser 20 vezes maior que a entrada (o Quick Sort "
        f"serial levou {ms(qs_k_grande_s)} nesse cenário). Na versão paralela, cada thread precisaria de um "
        "histograma próprio de tamanho k, multiplicando a memória alocada; por isso a implementação limita o "
        "número de histogramas a 2n/k e, quando esse limite é 1, realiza a contagem sequencialmente. Assim, com "
        f"k = 10.000.000 a versão de 8 threads registrou {ms(cs_k_grande_p)}, próximo da serial, evitando a "
        "degradação causada pela alocação de vários histogramas grandes."))
    B.append(("fig", os.path.join(DIR_FIG, "fig4_impacto_natureza.png"),
              "Impacto da natureza dos dados no tempo médio de execução (N = 100.000)"))
    B.append(("fig", os.path.join(DIR_FIG, "fig5_impacto_intervalo_valores.png"),
              "Tempo de execução em função do intervalo de valores [0..k] (N = 500.000)"))
    cab_int = ["Algoritmo", "Modo", "Threads", "k", "Média (ms)", "Desvio (ms)", "Speedup"]
    B.append(("tab", "Impacto do intervalo de valores [0..k] (N = 500.000, 5 amostras)", cab_int,
              d.linhas_intervalo()))

    B.append(("sub", "Variabilidade das amostras e gráficos dinâmicos"))
    cv_med = (d.resumo.DesvioPadraoMs / d.resumo.MediaMs)
    cv_1m = cv_med[d.resumo.Tamanho == 1_000_000].median()
    B.append(("par",
        "A Figura 6 apresenta a dispersão das cinco amostras para N = 1.000.000. O coeficiente de variação "
        f"mediano nesse tamanho foi de {pct(cv_1m)}, indicando medições estáveis após o aquecimento da JVM; a "
        "variabilidade é maior nas cargas pequenas, em que interferências do sistema operacional e do coletor de "
        "lixo têm peso relativo maior. Como ponto extra, foi desenvolvida uma aplicação de gráficos dinâmicos em "
        "Java Swing (Figura 7), que lê os CSVs ou executa novos testes e permite filtrar em tempo real por "
        "natureza dos dados, tamanho, número de threads e escala logarítmica, com dicas interativas sobre cada "
        "ponto medido."))
    B.append(("fig", os.path.join(DIR_FIG, "fig6_estabilidade_amostras_boxplot.png"),
              "Dispersão das 5 amostras para N = 1.000.000"))
    B.append(("fig", "resultados/graficos_swing/swing_tempo_por_tamanho.png",
              "Gráfico dinâmico gerado pela aplicação em Java Swing"))

    # --- Conclusão -------------------------------------------------------
    B.append(("secao", "Conclusão"))
    B.append(("par",
        "Os experimentos mostraram que o ganho com paralelismo depende tanto do algoritmo quanto do volume de "
        "dados. Para grandes volumes (N ≥ 100.000), as versões paralelas do Merge Sort e do Quick Sort com "
        f"Fork/Join oferecem o melhor equilíbrio entre tempo absoluto e escalabilidade, com Speedups de "
        f"{sp(s_ms_1m)} e {sp(s_qs_1m)}, embora a eficiência caia à medida que se adicionam threads. O Counting "
        "Sort é o mais rápido quando k é pequeno em relação a N, mas, sendo limitado pela memória, quase não "
        "escala com mais núcleos e degrada quando k cresce. O Bubble Sort Odd-Even confirma que um algoritmo "
        "quadrático pode ser acelerado pela divisão do trabalho de cada fase, mas o grande número de barreiras "
        "limita o Speedup. Por fim, para cargas pequenas o custo de coordenação das threads supera o ganho, o "
        "que reforça o uso de limiares de granularidade e a escolha do algoritmo adequado ao perfil dos dados "
        "antes de simplesmente aumentar o número de threads."))

    # --- Referências -----------------------------------------------------
    B.append(("secao", "Referências"))
    acesso = f"{HOJE.day} {MESES[HOJE.month - 1]} {ANO}"
    refs = [
        "AMDAHL, G. M. Validity of the single processor approach to achieving large scale computing "
        "capabilities. In: **AFIPS Spring Joint Computer Conference**, v. 30, p. 483-485, 1967.",
        "CORMEN, T. H.; LEISERSON, C. E.; RIVEST, R. L.; STEIN, C. **Algoritmos**: teoria e prática. "
        "4. ed. Rio de Janeiro: GEN LTC, 2024.",
        "GOETZ, B. et al. **Java Concurrency in Practice**. Boston: Addison-Wesley, 2006.",
        "GRAMA, A.; GUPTA, A.; KARYPIS, G.; KUMAR, V. **Introduction to Parallel Computing**. 2. ed. "
        "Harlow: Addison-Wesley, 2003.",
        "LEA, D. A Java Fork/Join Framework. In: **Proceedings of the ACM 2000 Conference on Java Grande**. "
        "New York: ACM, 2000. p. 36-43.",
        "ORACLE. **Java Platform, Standard Edition API Specification**: java.util.concurrent. Disponível em: "
        f"https://docs.oracle.com/en/java/javase/. Acesso em: {acesso}.",
        "SEDGEWICK, R.; WAYNE, K. **Algorithms**. 4. ed. Boston: Addison-Wesley, 2011.",
    ]
    for r in refs:
        B.append(("ref", r))

    # --- Anexos ----------------------------------------------------------
    B.append(("secao", "Anexos"))
    B.append(("link", "Repositório com o código-fonte completo, CSVs, aplicação Swing e Jupyter Notebook: ",
              URL_GITHUB))
    anexos = [
        ("Anexo A — Quick Sort (src/algoritmos/QuickSort.java)", "src/algoritmos/QuickSort.java"),
        ("Anexo B — Merge Sort (src/algoritmos/MergeSort.java)", "src/algoritmos/MergeSort.java"),
        ("Anexo C — Counting Sort (src/algoritmos/CountingSort.java)", "src/algoritmos/CountingSort.java"),
        ("Anexo D — Bubble Sort Odd-Even (src/algoritmos/BubbleSort.java)", "src/algoritmos/BubbleSort.java"),
        ("Anexo E — Framework de testes (src/benchmark/ExecutorBenchmark.java)",
         "src/benchmark/ExecutorBenchmark.java"),
    ]
    for titulo, caminho in anexos:
        B.append(("codigo", titulo, caminho))

    # Aviso de dados incoerentes (ex.: CSV gerado com a versão antiga do Bubble Sort)
    avisos = []
    for alg, tam in [("Bubble Sort", 50_000), ("Merge Sort", 1_000_000), ("Quick Sort", 1_000_000)]:
        f = d.resumo[(d.resumo.Algoritmo == alg) & (d.resumo.Tamanho == tam) & (d.resumo.Threads > 1)]
        # Eficiência muito acima de 100% indica CSV gerado pela versão antiga (Bubble em blocos + merge)
        if (f.Eficiencia > 3.0).any():
            avisos.append(f"{alg} (N = {tam}) apresenta eficiência acima de 300%")
    if not d.amb:
        avisos.append("resultados/ambiente.txt não encontrado (rode o benchmark com --cli)")
    return B, avisos


# ---------------------------------------------------------------------------
# Renderização DOCX sobre o modelo
# ---------------------------------------------------------------------------
RPR_ARIAL = '<w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial" w:eastAsia="Arial"/>'


class RenderDocx:
    def __init__(self, caminho_template):
        self.doc = Document(caminho_template)
        self.body = self.doc.element.body
        self.n_fig = 0
        self.n_tab = 0
        self._limpar_template()

    def _limpar_template(self):
        """Mantém apenas a tabela de cabeçalho (logo + UNIFOR) e a seção final do modelo."""
        primeira_tabela_vista = False
        for el in list(self.body):
            tag = el.tag.split("}")[1]
            if tag == "tbl" and not primeira_tabela_vista:
                primeira_tabela_vista = True
                continue
            if tag == "sectPr":
                continue
            self.body.remove(el)
        # Seção final: mantém margens do modelo, mas como página nova normal
        sect = self.body.find(qn("w:sectPr"))
        tipo = sect.find(qn("w:type"))
        if tipo is not None:
            sect.remove(tipo)

    def _anexar(self, el):
        self.body.insert(len(self.body) - 1, el)
        return el

    def _p(self, ppr, runs_xml):
        return self._anexar(parse_xml(f'<w:p {nsdecls("w")}><w:pPr>{ppr}</w:pPr>{runs_xml}</w:p>'))

    @staticmethod
    def _run(texto, tam=22, negrito=False, italico=False, cor=None, fonte=RPR_ARIAL, preservar=True):
        from xml.sax.saxutils import escape
        rpr = fonte
        if negrito:
            rpr += "<w:b/><w:bCs/>"
        if italico:
            rpr += "<w:i/><w:iCs/>"
        if cor:
            rpr += f'<w:color w:val="{cor}"/>'
        rpr += f'<w:sz w:val="{tam}"/><w:szCs w:val="{tam}"/><w:lang w:val="pt-BR"/>'
        esp = ' xml:space="preserve"' if preservar else ""
        return f"<w:r><w:rPr>{rpr}</w:rPr><w:t{esp}>{escape(texto)}</w:t></w:r>"

    def _runs_markup(self, texto, tam=22):
        """Converte trechos entre ** ** em negrito."""
        partes = texto.split("**")
        return "".join(self._run(p, tam, negrito=(i % 2 == 1)) for i, p in enumerate(partes) if p)

    # Blocos -------------------------------------------------------------
    def titulo(self, texto):
        self._p('<w:pStyle w:val="TAMainText"/><w:spacing w:before="360" w:after="240" w:line="360" '
                'w:lineRule="auto"/><w:ind w:firstLine="0"/><w:jc w:val="center"/>',
                self._run(texto.upper(), 28, negrito=True))

    def autor(self, texto):
        self._p('<w:pStyle w:val="TAMainText"/><w:spacing w:line="360" w:lineRule="auto"/>'
                '<w:ind w:firstLine="0"/><w:jc w:val="left"/>',
                self._run(texto, 22, italico=True))

    def palavras(self, texto):
        rotulo, resto = texto.split(":", 1)
        self._p('<w:pStyle w:val="NormalWeb"/><w:spacing w:before="280" w:after="240"/>',
                self._run(rotulo + ":", 18, negrito=True) + self._run(resto, 18))

    def secao(self, texto):
        # Mesma caixa cinza com borda do modelo (estilo Absbox)
        self._p('<w:pStyle w:val="Absbox"/><w:keepNext/><w:pBdr>'
                '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
                '<w:left w:val="single" w:sz="4" w:space="4" w:color="000000"/>'
                '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
                '<w:right w:val="single" w:sz="4" w:space="4" w:color="000000"/></w:pBdr>'
                '<w:shd w:val="clear" w:color="auto" w:fill="D9D9D9"/>'
                '<w:spacing w:before="240" w:after="120"/><w:ind w:left="85"/><w:jc w:val="left"/>',
                self._run(texto, 24, cor="000000"))

    def sub(self, texto):
        self._p('<w:pStyle w:val="TAMainText"/><w:keepNext/><w:spacing w:before="120" w:after="60" '
                'w:line="360" w:lineRule="auto"/><w:ind w:left="85" w:right="130" w:firstLine="0"/>'
                '<w:jc w:val="left"/>', self._run(texto, 22, negrito=True))

    def par(self, texto):
        self._p('<w:pStyle w:val="TAMainText"/><w:spacing w:after="120" w:line="360" w:lineRule="auto"/>'
                '<w:ind w:left="85" w:right="130" w:firstLine="709"/><w:jc w:val="both"/>',
                self._runs_markup(texto))

    def ref(self, texto):
        self._p('<w:pStyle w:val="TAMainText"/><w:spacing w:after="160" w:line="240" w:lineRule="auto"/>'
                '<w:ind w:left="85" w:right="130" w:firstLine="0"/><w:jc w:val="left"/>',
                self._runs_markup(texto))

    def _legenda(self, texto, depois=False):
        espaco = '<w:spacing w:before="0" w:after="200"/>' if depois else \
            '<w:spacing w:before="160" w:after="60"/><w:keepNext/>'
        self._p(f'<w:pStyle w:val="TAMainText"/>{espaco}<w:ind w:firstLine="0"/><w:jc w:val="center"/>',
                self._run(texto, 20 if not depois else 18, negrito=not depois))

    def fig(self, caminho, legenda):
        self.n_fig += 1
        self._legenda(f"Figura {self.n_fig} – {legenda}")
        if os.path.exists(caminho):
            p = self.doc.add_paragraph()
            self.body.remove(p._p)
            self._anexar(p._p)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(caminho, width=Cm(16.5))
        else:
            self.par(f"[Imagem não encontrada: {caminho}]")
        self._legenda(FONTE, depois=True)

    def tab(self, titulo, cab, linhas):
        self.n_tab += 1
        self._legenda(f"Tabela {self.n_tab} – {titulo}")
        t = self.doc.add_table(rows=1 + len(linhas), cols=len(cab))
        t.style = self.doc.styles["Table Grid"]
        self.body.remove(t._tbl)
        self._anexar(t._tbl)
        for j, h in enumerate(cab):
            self._celula(t.cell(0, j), h, cabecalho=True)
        for i, linha in enumerate(linhas, start=1):
            for j, v in enumerate(linha):
                self._celula(t.cell(i, j), v, alinhar_dir=(j >= 2))
        # repete cabeçalho em quebras de página
        trpr = t.rows[0]._tr.get_or_add_trPr()
        trpr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
        self._legenda(FONTE, depois=True)

    def _celula(self, cel, texto, cabecalho=False, alinhar_dir=False):
        cel.text = ""
        p = cel.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.first_line_indent = Cm(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (cabecalho or alinhar_dir) else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(texto)
        r.font.name = "Arial"
        r.font.size = Pt(8.5)
        r.bold = cabecalho
        if cabecalho:
            cel._tc.get_or_add_tcPr().append(
                parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="D9D9D9"/>'))

    def link(self, prefixo, url):
        p = self._p('<w:pStyle w:val="TAMainText"/><w:spacing w:after="160" w:line="360" w:lineRule="auto"/>'
                    '<w:ind w:left="85" w:right="130" w:firstLine="0"/>', self._run(prefixo, 22))
        rid = self.doc.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
                                           "hyperlink", is_external=True)
        hl = parse_xml(f'<w:hyperlink {nsdecls("w", "r")} r:id="{rid}">'
                       f'{self._run(url, 22, cor="0563C1")}</w:hyperlink>')
        hl.find(qn("w:r")).find(qn("w:rPr")).append(parse_xml(f'<w:u {nsdecls("w")} w:val="single"/>'))
        p.append(hl)

    def codigo(self, titulo, caminho):
        self.sub(titulo)
        if not os.path.exists(caminho):
            self.par(f"[Arquivo não encontrado: {caminho}]")
            return
        with open(caminho, encoding="utf-8") as f:
            linhas = f.read().replace("\t", "    ").rstrip("\n").split("\n")
        fonte = '<w:rFonts w:ascii="Courier New" w:hAnsi="Courier New" w:cs="Courier New"/>'
        runs = []
        for i, linha in enumerate(linhas):
            if i:
                runs.append("<w:r><w:br/></w:r>")
            runs.append(self._run(linha if linha else " ", 15, fonte=fonte))
        self._p('<w:pBdr><w:top w:val="single" w:sz="4" w:space="4" w:color="BFBFBF"/>'
                '<w:left w:val="single" w:sz="4" w:space="4" w:color="BFBFBF"/>'
                '<w:bottom w:val="single" w:sz="4" w:space="4" w:color="BFBFBF"/>'
                '<w:right w:val="single" w:sz="4" w:space="4" w:color="BFBFBF"/></w:pBdr>'
                '<w:shd w:val="clear" w:color="auto" w:fill="F2F2F2"/>'
                '<w:spacing w:before="60" w:after="200" w:line="240" w:lineRule="auto"/>'
                '<w:ind w:left="85" w:right="130"/><w:jc w:val="left"/>', "".join(runs))

    def aviso(self, texto):
        self._p('<w:shd w:val="clear" w:color="auto" w:fill="FFE0E0"/><w:spacing w:after="120"/>',
                self._run(texto, 20, negrito=True, cor="C00000"))

    def renderizar(self, blocos, avisos):
        for b in blocos:
            getattr(self, b[0])(*b[1:])
            if b[0] == "palavras" and avisos:
                self.aviso("ATENÇÃO (remova antes de entregar): os dados parecem desatualizados — "
                           + "; ".join(avisos) + ". Rode novamente o benchmark e os scripts.")

    def salvar(self, caminho):
        self.doc.core_properties.title = TITULO
        self.doc.core_properties.author = f"{AUTOR_1}; {AUTOR_2}"
        self.doc.save(caminho)


# ---------------------------------------------------------------------------
# Renderização Markdown
# ---------------------------------------------------------------------------
def renderizar_markdown(blocos):
    md = ["| ![Unifor](templates/unifor_logo.png) | **UNIVERSIDADE DE FORTALEZA**<br>**CENTRO DE CIÊNCIAS "
          "TECNOLÓGICAS**<br>**CURSO: CIÊNCIA DA COMPUTAÇÃO** |", "| :---: | :--- |", ""]
    n_fig = n_tab = 0
    for b in blocos:
        tipo = b[0]
        if tipo == "titulo":
            md += [f"# {b[1].upper()}", ""]
        elif tipo == "autor":
            md += [f"*{b[1]}*  "]
        elif tipo == "palavras":
            rot, resto = b[1].split(":", 1)
            md += ["", f"**{rot}:**{resto}", ""]
        elif tipo == "secao":
            md += ["", f"## {b[1]}", ""]
        elif tipo == "sub":
            md += [f"### {b[1]}", ""]
        elif tipo in ("par", "ref"):
            md += [b[1], ""]
        elif tipo == "fig":
            n_fig += 1
            md += [f"**Figura {n_fig} – {b[2]}**", "", f"![Figura {n_fig}]({b[1]})", "", f"*{FONTE}*", ""]
        elif tipo == "tab":
            n_tab += 1
            _, titulo, cab, linhas = b
            md += [f"**Tabela {n_tab} – {titulo}**", "", "| " + " | ".join(cab) + " |",
                   "|" + "---|" * len(cab)]
            md += ["| " + " | ".join(l) + " |" for l in linhas]
            md += ["", f"*{FONTE}*", ""]
        elif tipo == "link":
            md += [f"{b[1]}[{b[2]}]({b[2]})", ""]
        elif tipo == "codigo":
            md += [f"- [{b[1]}]({b[2]})"]
    return "\n".join(md) + "\n"


def extrair_logo_template():
    """Salva o logo oficial do modelo para uso no Markdown."""
    import zipfile
    destino = "templates/unifor_logo.png"
    if not os.path.exists(destino):
        with zipfile.ZipFile(TEMPLATE) as z:
            with z.open("word/media/image1.png") as src, open(destino, "wb") as dst:
                dst.write(src.read())


def main():
    if not os.path.exists(TEMPLATE):
        sys.exit(f"Modelo não encontrado: {TEMPLATE}")
    dados = Dados()
    blocos, avisos = montar_conteudo(dados)

    r = RenderDocx(TEMPLATE)
    r.renderizar(blocos, avisos)
    r.salvar(SAIDA_DOCX)
    print("Relatório Word gerado:", SAIDA_DOCX)

    extrair_logo_template()
    with open(SAIDA_MD, "w", encoding="utf-8") as f:
        f.write(renderizar_markdown(blocos))
    print("Relatório Markdown gerado:", SAIDA_MD)

    for a in avisos:
        print("AVISO:", a)


if __name__ == "__main__":
    main()
